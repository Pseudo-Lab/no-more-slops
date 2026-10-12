#!/usr/bin/env python3
"""Supplied visual materials (figures, diagrams, tables) for paper-to-latex.

Subcommands
-----------
inventory  <paper_file> <project_dir> <path> [<path> ...]
    Inventory every supplied file/directory, extract metadata (format, vector or
    raster, size/DPI, embedded text, table cells, sidecar captions), find the
    manuscript's figure/table "slots" (captions, mentions, image paths,
    placeholders), propose a slot for each asset with a score, and flag
    ambiguity, conflicts, unmatched assets, slots without assets, and
    discrepancies between supplied tables and the manuscript.
    Writes <project_dir>/_conversion/visuals_manifest.json. Each asset has a
    "decision" object that the converter (Claude) must fill in before `check`.

convert    <asset_file> <dest_dir> [--name <stem>]
    Make a LaTeX-compatible copy: PDF/PNG/JPG copied byte-for-byte; EPS -> PDF
    (epstopdf); SVG -> PDF (inkscape | rsvg-convert | cairosvg | svglib);
    TIFF/GIF/BMP/WebP -> PNG (Pillow, lossless). Verifies conversions
    (aspect ratio, text labels preserved). Prints a JSON record. Exit 3 when no
    converter is available. Never uses Windows' convert.exe.

table      <table_file> <out_tex> [--caption-file F | --caption TEXT] [--label L]
           [--wide auto|yes|no] [--placement t]
    CSV/TSV/Markdown pipe table -> booktabs LaTeX table with every cell copied
    verbatim (only LaTeX escaping / Unicode-math mapping). No caption is
    invented: without --caption/--caption-file the table gets no \\caption.
    A .tex table is copied unchanged (wrap it by hand if it lacks a float).

check      <project_dir> [--manifest M] [--render-dpi 110]
    After compiling: every asset has a decision; included assets exist, are
    referenced, unchanged (sha256 or verified conversion), captioned with the
    recorded verbatim caption, labelled and cross-referenced, placed inside the
    main paper (unless decided otherwise), not trimmed/clipped/distorted, and
    legible (effective text size of PDF figures, effective DPI of rasters).
    Table numbers in the project file must equal the supplied table's numbers.
    Renders each page holding a visual plus each source asset into
    _conversion/render/ for visual inspection. Writes
    _conversion/visuals_check.json. Exit 1 on any ERROR.
"""
import csv
import hashlib
import html
import json
import os
import re
import shutil
import struct
import subprocess
import sys
from collections import Counter
from pathlib import Path

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:  # noqa: BLE001
        pass

FIG_EXT = {".pdf", ".svg", ".png", ".jpg", ".jpeg", ".eps", ".tif", ".tiff", ".gif", ".bmp", ".webp"}
TABLE_EXT = {".csv", ".tsv", ".md", ".markdown", ".tex", ".xlsx", ".xls"}
CAPTION_SIDECAR = re.compile(r"\.(caption|cap)\.(txt|md|tex)$", re.I)
EVIDENCE_EXT = {".py", ".ipynb", ".r", ".json", ".yaml", ".yml", ".log", ".txt", ".mmd", ".dot", ".gv"}
SKIP_DIRS = {"_build", "_conversion", "__pycache__", ".git", ".ipynb_checkpoints"}
VECTOR_PREF = [".pdf", ".eps", ".svg", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".gif", ".bmp", ".webp"]
NUM_RE = re.compile(r"(?<![\w.])[-\u2212]?\d[\d,]*(?:\.\d+)?%?")
STOP = {"the", "a", "an", "of", "and", "or", "for", "in", "on", "to", "with", "by", "vs", "versus", "from",
        "fig", "figure", "table", "tab", "final", "v1", "v2", "v3", "new", "copy", "img", "image", "plot",
        "chart", "graph", "is", "are", "at", "as", "per", "our"}


# ---------------------------------------------------------------- utilities
def sha256(p):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def run(cmd, timeout=120):
    try:
        r = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace",
                           timeout=timeout)
        return r.returncode, r.stdout, r.stderr
    except (OSError, subprocess.TimeoutExpired) as exc:
        return 1, "", str(exc)


def norm_num(tok):
    return tok.replace("\u2212", "-").replace(",", "")


def numbers(text):
    return [norm_num(m.group(0)) for m in NUM_RE.finditer(text or "")]


def alnum(text):
    return "".join(ch.lower() for ch in (text or "") if ch.isalnum())


def tokens(text):
    text = re.sub(r"([a-z])([A-Z])", r"\1 \2", text or "")
    return {t for t in re.findall(r"[a-z][a-z0-9]+", text.lower()) if t not in STOP and len(t) > 2}


def tex_to_plain(s):
    """Rough TeX -> text for comparisons (keeps numbers and words)."""
    s = re.sub(r"(?<!\\)%.*", "", s)
    s = s.replace("$-$", "-").replace("\\textminus", "-").replace("--", "-")
    s = re.sub(r"\\(?:textbf|textit|emph|underline|mathbf|mathrm|text|textsc|texttt)\s*\{", "{", s)
    s = re.sub(r"\\(pm|times|le|ge|leq|geq|Delta|alpha|beta|gamma|sigma|mu|cdot)\b", lambda m: " " + m.group(1) + " ", s)
    s = re.sub(r"\\[a-zA-Z]+\*?", " ", s)
    s = re.sub(r"[{}$~]", " ", s)
    s = s.replace("\\&", "&").replace("\\%", "%").replace("\\_", "_").replace("\\#", "#")
    return " ".join(s.split())


def brace_arg(text, start):
    """Return (content, end_index) of the {...} group starting at text[start] == '{'."""
    if start >= len(text) or text[start] != "{":
        return None, start
    depth, i = 0, start
    while i < len(text):
        c = text[i]
        if c == "\\":
            i += 2
            continue
        if c == "{":
            depth += 1
        elif c == "}":
            depth -= 1
            if depth == 0:
                return text[start + 1:i], i + 1
        i += 1
    return None, start


def safe_stem(stem):
    s = re.sub(r"[^A-Za-z0-9_-]+", "_", stem).strip("_")
    return s or "visual"


# ---------------------------------------------------------- asset metadata
def png_info(p):
    with open(p, "rb") as f:
        head = f.read(64 * 1024)
    if head[:8] != b"\x89PNG\r\n\x1a\n":
        return {}
    w, h = struct.unpack(">II", head[16:24])
    info = {"px_w": w, "px_h": h}
    i = head.find(b"pHYs")
    if i > 0:
        xppu, yppu, unit = struct.unpack(">IIB", head[i + 4:i + 13])
        if unit == 1 and xppu:
            info["dpi"] = round(xppu * 0.0254)
    return info


def raster_info(p):
    info = {}
    try:
        from PIL import Image  # type: ignore
        with Image.open(p) as im:
            info = {"px_w": im.width, "px_h": im.height, "mode": im.mode}
            dpi = im.info.get("dpi")
            if dpi:
                info["dpi"] = round(float(dpi[0]))
    except Exception:  # noqa: BLE001
        if p.suffix.lower() == ".png":
            info = png_info(p)
    return info


def pdf_info(p):
    info = {}
    rc, out, _ = run(["pdfinfo", str(p)])
    if rc == 0:
        m = re.search(r"Pages:\s+(\d+)", out)
        info["pages"] = int(m.group(1)) if m else None
        m = re.search(r"Page size:\s+([\d.]+) x ([\d.]+) pts", out)
        if m:
            info["width_pt"], info["height_pt"] = float(m.group(1)), float(m.group(2))
    rc, out, _ = run(["pdftotext", "-l", "1", str(p), "-"])
    info["text"] = " ".join(out.split()) if rc == 0 else ""
    hs = text_heights_pdf(p)
    if hs:
        hs.sort()
        info["text_height_min_pt"] = round(hs[0], 2)
        info["text_height_p10_pt"] = round(hs[max(0, len(hs) // 10)], 2)
    rc, out, _ = run(["pdffonts", "-l", "1", str(p)])
    if rc == 0:
        not_emb = [l.split()[0] for l in out.splitlines()[2:] if l.strip() and re.search(r"\s(no)\s+(yes|no)\s+(yes|no)\s+\d+\s+\d+\s*$", l)]
        info["fonts_not_embedded"] = not_emb
    return info


def text_heights_pdf(p):
    """Effective font sizes (pt) of text on page 1. pypdf first (font size x text/CTM scale);
    fallback: poppler `pdftotext -bbox` word heights (xpdf builds lack -bbox)."""
    try:
        from pypdf import PdfReader  # type: ignore
        sizes = []

        def visit(text, cm, tm, font_dict, font_size):
            if not alnum(text) or not font_size:
                return
            a = tm[0] * cm[0] + tm[1] * cm[2]
            b = tm[0] * cm[1] + tm[1] * cm[3]
            c = tm[2] * cm[0] + tm[3] * cm[2]
            d = tm[2] * cm[1] + tm[3] * cm[3]
            scale = abs(a * d - b * c) ** 0.5
            if scale:
                sizes.append(font_size * scale)

        PdfReader(str(p)).pages[0].extract_text(visitor_text=visit)
        if sizes:
            return sizes
    except Exception:  # noqa: BLE001
        pass
    rc, out, _ = run(["pdftotext", "-bbox", "-l", "1", str(p), "-"])
    if rc != 0:
        return []
    hs = []
    for m in re.finditer(r'<word xMin="([\d.]+)" yMin="([\d.]+)" xMax="([\d.]+)" yMax="([\d.]+)">([^<]*)</word>', out):
        h = float(m.group(4)) - float(m.group(2))
        if h > 0.5 and alnum(html.unescape(m.group(5))):
            hs.append(h)
    return hs


def svg_info(p):
    txt = p.read_text(encoding="utf-8", errors="replace")
    info = {}
    m = re.search(r"<svg\b[^>]*>", txt, re.S)
    tag = m.group(0) if m else ""
    for k in ("width", "height", "viewBox"):
        mm = re.search(rf'\b{k}="([^"]+)"', tag)
        if mm:
            info[k] = mm.group(1)
    texts = [html.unescape(re.sub(r"<[^>]+>", " ", t)) for t in re.findall(r"<text\b[^>]*>(.*?)</text>", txt, re.S)]
    info["text"] = " ".join(" ".join(texts).split())
    info["text_as_paths"] = not texts and "<path" in txt
    info["has_foreign_object"] = "<foreignObject" in txt
    info["external_refs"] = re.findall(r'(?:xlink:)?href="(?!#|data:)([^"]+)"', txt)
    return info


def eps_info(p):
    head = p.read_bytes()[:4096].decode("latin-1", "replace")
    m = re.search(r"%%BoundingBox:\s*(-?\d+)\s+(-?\d+)\s+(-?\d+)\s+(-?\d+)", head)
    return {"bbox": [int(x) for x in m.groups()]} if m else {}


# ------------------------------------------------------------ table parsing
def split_md_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|") and not line.endswith("\\|"):
        line = line[:-1]
    return [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", line)]


MD_SEP = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?\s*$")
CAP_LINE = re.compile(r"^\s*(?:\*\*|__)?\s*(Table|Figure|Fig\.)\s*(\d+)\s*[.:]?\s*(?:\*\*|__)?\s*[.:]?\s*(.*?)\s*$", re.I)


def parse_md_tables(text):
    """Return list of {rows, align, caption, caption_num, start_line}."""
    lines = text.splitlines()
    out, i = [], 0
    while i < len(lines):
        if "|" in lines[i] and i + 1 < len(lines) and MD_SEP.match(lines[i + 1]):
            header = split_md_row(lines[i])
            align = []
            for c in split_md_row(lines[i + 1]):
                c = c.strip()
                align.append("c" if c.startswith(":") and c.endswith(":") else "r" if c.endswith(":") else "l")
            rows = [header]
            j = i + 2
            while j < len(lines) and "|" in lines[j] and lines[j].strip():
                rows.append(split_md_row(lines[j]))
                j += 1
            cap, num = None, None
            for k in list(range(i - 1, max(-1, i - 4), -1)) + list(range(j, min(len(lines), j + 3))):
                m = CAP_LINE.match(lines[k]) if lines[k].strip() else None
                if m and m.group(1).lower() == "table":
                    num, cap = int(m.group(2)), m.group(3)
                    break
            out.append({"rows": rows, "align": align, "caption": cap, "caption_num": num, "start_line": i + 1})
            i = j
        else:
            i += 1
    return out


def parse_tex_tables(text):
    out = []
    for m in re.finditer(r"\\begin\{tabular\*?\}", text):
        start = m.end()
        # skip width arg for tabular* and the column spec
        pos = start
        if text[m.start():m.end()].startswith("\\begin{tabular*}"):
            _, pos = brace_arg(text, pos)
        spec, pos = brace_arg(text, pos)
        end = text.find("\\end{tabular", pos)
        body = text[pos:end]
        body = re.sub(r"\\(toprule|midrule|bottomrule|hline|cmidrule(\([^)]*\))?(\{[^}]*\})?)", "", body)
        rows = []
        for r in re.split(r"\\\\(?:\[[^\]]*\])?", body):
            if not r.strip():
                continue
            rows.append([c.strip() for c in re.split(r"(?<!\\)&", r)])
        out.append({"rows": rows, "colspec": spec})
    caps = []
    for m in re.finditer(r"\\caption(?:\[[^\]]*\])?\s*(?=\{)", text):
        c, _ = brace_arg(text, m.end())
        caps.append(c)
    labels = re.findall(r"\\label\{([^}]*)\}", text)
    return out, caps, labels


def tex_table_numbers(tex):
    """Numbers in the cells of every tabular in tex (ignores captions, options, lengths)."""
    ts, _, _ = parse_tex_tables(tex)
    return numbers(tex_to_plain(" ".join(" ".join(" ".join(r) for r in t["rows"]) for t in ts)))


def read_table(p):
    ext = p.suffix.lower()
    res = {"rows": [], "caption": None, "labels": []}
    if ext in (".csv", ".tsv"):
        with open(p, newline="", encoding="utf-8-sig", errors="replace") as f:
            res["rows"] = [r for r in csv.reader(f, delimiter="\t" if ext == ".tsv" else ",") if any(c.strip() for c in r)]
    elif ext in (".md", ".markdown"):
        ts = parse_md_tables(p.read_text(encoding="utf-8", errors="replace"))
        if ts:
            res.update({"rows": ts[0]["rows"], "align": ts[0]["align"], "caption": ts[0]["caption"],
                        "caption_num": ts[0]["caption_num"], "n_tables_in_file": len(ts)})
    elif ext == ".tex":
        txt = p.read_text(encoding="utf-8", errors="replace")
        ts, caps, labels = parse_tex_tables(txt)
        if ts:
            res["rows"] = ts[0]["rows"]
            res["n_tables_in_file"] = len(ts)
        res["caption"] = caps[0] if caps else None
        res["labels"] = labels
        res["has_float_env"] = bool(re.search(r"\\begin\{table\*?\}", txt))
    elif ext in (".xlsx", ".xls"):
        try:
            import openpyxl  # type: ignore
            wb = openpyxl.load_workbook(p, data_only=False, read_only=True)
            ws = wb.worksheets[0]
            res["rows"] = [["" if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)]
            res["sheets"] = wb.sheetnames
            res["warning"] = "xlsx read with openpyxl: cell values are the stored values; check number formatting (displayed rounding) against the manuscript"
        except ImportError:
            res["unsupported"] = "xlsx needs openpyxl (not installed); ask the user for CSV export"
    flat = " ".join(" ".join(r) for r in res["rows"])
    res["numbers"] = numbers(tex_to_plain(flat) if ext == ".tex" else flat)
    res["shape"] = [len(res["rows"]), max((len(r) for r in res["rows"]), default=0)]
    return res


# ------------------------------------------------------- manuscript slots
def manuscript_slots(paper, project):
    """Find figure/table slots in the manuscript."""
    ext = paper.suffix.lower()
    plain_p = project / "_conversion" / "source_plain.txt"
    raw = paper.read_text(encoding="utf-8", errors="replace") if ext in (".md", ".markdown", ".tex") else \
        (plain_p.read_text(encoding="utf-8", errors="replace") if plain_p.exists() else "")
    slots = {}

    def slot(kind, num):
        key = f"{kind}:{num}"
        if key not in slots:
            slots[key] = {"key": key, "kind": kind, "number": num, "caption": None, "mentions": 0,
                          "source_image": None, "placeholder": None, "manuscript_table": None, "label": None}
        return slots[key]

    if ext == ".tex":
        for env in re.finditer(r"\\begin\{(figure|table)\*?\}(.*?)\\end\{\1\*?\}", raw, re.S):
            kind, body = env.group(1), env.group(2)
            n = sum(1 for k in slots if k.startswith(kind)) + 1
            s = slot(kind, n)
            mcap = re.search(r"\\caption(?:\[[^\]]*\])?\s*(?=\{)", body)
            if mcap:
                s["caption"], _ = brace_arg(body, mcap.end())
            mlab = re.search(r"\\label\{([^}]*)\}", body)
            s["label"] = mlab.group(1) if mlab else None
            mg = re.search(r"\\includegraphics(?:\[[^\]]*\])?\{([^}]*)\}", body)
            s["source_image"] = mg.group(1) if mg else None
            if kind == "table":
                ts, _, _ = parse_tex_tables(body)
                if ts:
                    s["manuscript_table"] = ts[0]["rows"]
            if s["label"]:
                s["mentions"] += len(re.findall(r"\\(?:ref|cref|Cref|autoref)\{" + re.escape(s["label"]) + r"\}", raw))
    else:
        lines = raw.splitlines()
        for i, line in enumerate(lines):
            m = CAP_LINE.match(line)
            if m and len(line) < 600 and (line.strip().startswith(("*", "_", "T", "F")) and re.match(r"^\s*(\*\*|__)?\s*(Table|Figure|Fig\.)\s*\d+\s*[.:]", line, re.I)):
                kind = "table" if m.group(1).lower() == "table" else "figure"
                s = slot(kind, int(m.group(2)))
                s["caption"] = s["caption"] or m.group(3)
                s["caption_line"] = i + 1
        for m in re.finditer(r"!\[([^\]]*)\]\(\s*(<[^>]+>|[^)\"]+?)\s*(?:\"[^\"]*\")?\s*\)", raw):
            alt, path = m.group(1), m.group(2).strip("<>")
            mm = re.match(r"\s*(?:Figure|Fig\.)\s*(\d+)\s*[.:]?\s*(.*)", alt, re.I)
            if mm:
                s = slot("figure", int(mm.group(1)))
                s["caption"] = s["caption"] or mm.group(2)
            else:
                n = sum(1 for k in slots if k.startswith("figure")) + 1
                s = slot("figure", n)
            s["source_image"] = path
        for t in parse_md_tables(raw):
            if t["caption_num"]:
                s = slot("table", t["caption_num"])
                s["manuscript_table"] = t["rows"]
                s["caption"] = s["caption"] or t["caption"]
        for m in re.finditer(r"\[(?:insert\s+|place\s+)?(figure|fig\.|table)\s*(\d+)[^\]\n]{0,60}\]", raw, re.I):
            kind = "table" if m.group(1).lower() == "table" else "figure"
            s = slot(kind, int(m.group(2)))
            s["placeholder"] = m.group(0)
    # in-text mentions: ignore caption lines, image syntax and TeX captions
    body = re.sub(r"!\[[^\]]*\]\([^)]*\)", " ", raw)
    body = re.sub(r"\[(?:insert\s+|place\s+)?(?:figure|fig\.|table)\s*\d+[^\]\n]{0,60}\]", " ", body, flags=re.I)
    body = "\n".join(l for l in body.splitlines()
                     if not re.match(r"^\s*(\*\*|__)?\s*(Table|Figure|Fig\.)\s*\d+\s*[.:]", l, re.I))
    body = re.sub(r"\\caption(?:\[[^\]]*\])?\{[^{}]*\}", " ", body)
    for m in re.finditer(r"\b(Figure|Fig\.|Table)\s*~?(\d+)", body):
        kind = "table" if m.group(1).lower() == "table" else "figure"
        slot(kind, int(m.group(2)))["mentions"] += 1
    return list(slots.values())


# ------------------------------------------------------------- inventory
def collect(paths):
    files = []
    for p in paths:
        p = Path(p)
        if p.is_dir():
            for q in sorted(p.rglob("*")):
                if q.is_file() and not any(part in SKIP_DIRS or part.startswith(".") for part in q.relative_to(p).parts):
                    files.append(q)
        elif p.is_file():
            files.append(p)
        else:
            files.append(p)  # reported as missing
    seen, out = set(), []
    for f in files:
        k = str(f.resolve()).lower()
        if k not in seen:
            seen.add(k)
            out.append(f)
    return out


def classify(p):
    ext = p.suffix.lower()
    name = p.name.lower()
    if CAPTION_SIDECAR.search(name):
        return "caption_sidecar"
    if ext in FIG_EXT:
        return "figure"
    if ext in (".csv", ".tsv", ".xlsx", ".xls"):
        return "table"
    if ext in (".md", ".markdown"):
        txt = p.read_text(encoding="utf-8", errors="replace")
        return "table" if parse_md_tables(txt) else "document"
    if ext == ".tex":
        txt = p.read_text(encoding="utf-8", errors="replace")
        if re.search(r"\\begin\{tabular", txt):
            return "table"
        if re.search(r"\\begin\{(tikzpicture|figure)", txt):
            return "figure_tex"
        return "document"
    if ext in EVIDENCE_EXT:
        return "evidence"
    return "unsupported"


def figure_kind_hint(stem, text):
    t = (stem + " " + (text or "")).lower()
    if re.search(r"workflow|pipeline|overview|architecture|method|framework|diagram|flow|schematic|procedure", t):
        return "diagram"
    if re.search(r"plot|curve|bar|scatter|hist|box|violin|heatmap|roc|accuracy|loss|f1|auc|score|ci|error", t):
        return "statistical plot"
    return "figure"


_MS_NUMS = set()   # numbers quoted anywhere in the manuscript (filled by inventory)


def score(asset, s):
    sc, why = 0, []
    if asset["kind"] == "table" and s["kind"] != "table":
        return 0, ["kind mismatch"]
    if asset["kind"] in ("figure", "figure_tex") and s["kind"] != "figure":
        if not (asset.get("table_as_image") and s["kind"] == "table"):
            return 0, ["kind mismatch"]
    if s.get("source_image"):
        src_stem = safe_stem(Path(s["source_image"]).stem).lower()
        if src_stem == asset["group"].lower():
            sc += 100
            why.append("manuscript image path matches file name")
    if asset.get("file_number") == s["number"] and asset.get("file_kind") in (None, s["kind"]):
        sc += 50
        why.append(f"file name number {s['number']}")
    if asset["kind"] == "table":   # header row describes the table; body cells only dilute the overlap
        desc = " ".join((asset.get("rows") or [[]])[0])
    else:
        desc = " ".join((asset.get("text") or "").split()[:80])
    a_tok = tokens(asset["group"]) | tokens(asset.get("sidecar_caption") or "") | \
        tokens(desc) | tokens(asset.get("caption_in_file") or "")
    s_tok = tokens(s.get("caption") or "") | tokens(s.get("label") or "")
    if a_tok and s_tok:
        inter = a_tok & s_tok
        ov = len(inter) / max(2, min(len(a_tok), len(s_tok)))
        if inter:
            sc += round(40 * min(1.0, ov))
            why.append("caption/name overlap: " + ",".join(sorted(inter)[:6]))
    if asset["kind"] == "table" and s.get("manuscript_table"):
        mn = Counter(numbers(" ".join(" ".join(r) for r in s["manuscript_table"])))
        an = Counter(asset.get("numbers") or [])
        if mn and an:
            ratio = sum((mn & an).values()) / max(sum(mn.values()), sum(an.values()))
            sc += round(60 * ratio)
            why.append(f"numeric overlap {ratio:.2f}")
    if asset["kind"] == "table" and s["kind"] == "table" and not s.get("manuscript_table") and _MS_NUMS:
        dec = [n for n in (asset.get("numbers") or []) if "." in n]
        if dec:
            quoted = sum(1 for n in set(dec) if n in _MS_NUMS) / len(set(dec))
            if quoted:
                sc += round(30 * quoted)
                why.append(f"{quoted:.0%} of its decimal values are quoted in the manuscript text")
    return sc, why


def table_discrepancy(asset_rows, ms_rows):
    a = [[" ".join(tex_to_plain(c).split()) for c in r] for r in asset_rows]
    m = [[" ".join(tex_to_plain(c).replace("**", "").split()) for c in r] for r in ms_rows]
    diffs = []
    if len(a) == len(m) and all(len(x) == len(y) for x, y in zip(a, m)):
        for i, (ra, rm) in enumerate(zip(a, m)):
            for j, (ca, cm) in enumerate(zip(ra, rm)):
                if numbers(ca) != numbers(cm) or (not numbers(ca) and alnum(ca) != alnum(cm)):
                    diffs.append({"row": i, "col": j, "manuscript": rm[j], "supplied": ra[j]})
        return {"comparable": "cell-by-cell", "differences": diffs}
    na, nm = Counter(numbers(" ".join(map(" ".join, a)))), Counter(numbers(" ".join(map(" ".join, m))))
    return {"comparable": "numbers only (shapes differ: manuscript %dx%d, supplied %dx%d)" % (
        len(m), max(map(len, m), default=0), len(a), max(map(len, a), default=0)),
        "only_in_manuscript": sorted((nm - na).elements()), "only_in_supplied": sorted((na - nm).elements())}


def cmd_inventory(args):
    if len(args) < 3:
        sys.exit("usage: visuals.py inventory <paper_file> <project_dir> <path> [<path> ...]")
    paper, project = Path(args[0]), Path(args[1])
    (project / "_conversion").mkdir(parents=True, exist_ok=True)
    files = collect(args[2:])
    issues, sidecars, others, assets = [], {}, [], []
    for f in files:
        if not f.exists():
            issues.append({"severity": "ERROR", "msg": f"supplied path does not exist: {f}"})
            continue
        if f.resolve() == paper.resolve():
            continue
        c = classify(f)
        if c == "caption_sidecar":
            sidecars[CAPTION_SIDECAR.sub("", f.name).lower()] = f
        elif c in ("figure", "table", "figure_tex"):
            assets.append({"path": str(f.resolve()), "kind": c, "ext": f.suffix.lower()})
        else:
            others.append({"path": str(f.resolve()), "class": c})

    # group variants: same stem, different formats (fig1.pdf + fig1.png + fig1.svg)
    groups = {}
    for a in assets:
        stem = safe_stem(Path(a["path"]).stem)
        groups.setdefault((a["kind"] if a["kind"] != "figure_tex" else "figure", stem.lower()), []).append(a)
    merged = []
    for (kind, _), variants in groups.items():
        variants.sort(key=lambda v: VECTOR_PREF.index(v["ext"]) if v["ext"] in VECTOR_PREF else 99)
        if kind == "table":
            order = {".tex": 0, ".csv": 1, ".tsv": 2, ".md": 3, ".markdown": 3, ".xlsx": 4, ".xls": 5}
            variants.sort(key=lambda v: order.get(v["ext"], 9))
        primary = dict(variants[0])
        primary["group"] = safe_stem(Path(primary["path"]).stem)
        primary["variants"] = [v["path"] for v in variants]
        merged.append(primary)

    hashes = {}
    for a in merged:
        p = Path(a["path"])
        a["sha256"] = sha256(p)
        a["bytes"] = p.stat().st_size
        hashes.setdefault(a["sha256"], []).append(a["path"])
        m = re.search(r"(?:^|[^a-z])(fig(?:ure)?|tab(?:le)?)[\s_.-]*0*(\d+)", Path(a["path"]).stem.lower())
        if m:
            a["file_kind"] = "figure" if m.group(1).startswith("fig") else "table"
            a["file_number"] = int(m.group(2))
        sc = sidecars.get(Path(a["path"]).stem.lower()) or sidecars.get(Path(a["path"]).name.lower())
        if sc:
            a["sidecar_caption_file"] = str(sc.resolve())
            a["sidecar_caption"] = " ".join(sc.read_text(encoding="utf-8", errors="replace").split())
        ext = a["ext"]
        if a["kind"] == "table":
            t = read_table(p)
            a.update({"rows": t["rows"], "numbers": t["numbers"], "shape": t["shape"],
                      "caption_in_file": t.get("caption"), "labels_in_file": t.get("labels", []),
                      "has_float_env": t.get("has_float_env")})
            for k in ("unsupported", "warning", "n_tables_in_file"):
                if t.get(k):
                    a[k] = t[k]
            a["text"] = " ".join(" ".join(r) for r in t["rows"])
        elif ext == ".pdf":
            a.update(pdf_info(p))
            a["vector"] = True
        elif ext == ".svg":
            a.update(svg_info(p))
            a["vector"] = True
            a["needs_conversion"] = "svg -> pdf"
        elif ext == ".eps":
            a.update(eps_info(p))
            a["vector"] = True
            a["needs_conversion"] = "eps -> pdf"
        elif a["kind"] == "figure":
            a.update(raster_info(p))
            a["vector"] = False
            if ext not in (".png", ".jpg", ".jpeg"):
                a["needs_conversion"] = f"{ext[1:]} -> png"
        if a["kind"] == "figure":
            a["figure_type"] = figure_kind_hint(a["group"], a.get("text"))
            if re.search(r"tab(le)?[\s_.-]*\d|results?_table", a["group"].lower()):
                a["table_as_image"] = True
        a["decision"] = {"action": None, "slot": None, "project_path": None, "label": None,
                         "caption_text": None, "caption_source": None, "environment": None,
                         "placement": "main", "reason": None}

    slots = manuscript_slots(paper, project)
    plain_p = project / "_conversion" / "source_plain.txt"
    ms_text = plain_p.read_text(encoding="utf-8", errors="replace") if plain_p.exists() else \
        paper.read_text(encoding="utf-8", errors="replace")
    _MS_NUMS.clear()
    _MS_NUMS.update(numbers(ms_text))
    # slots still needing a visual (no image / table body in the manuscript itself)
    open_slots = {k: [s for s in slots if s["kind"] == k and not s.get("manuscript_table")
                      and not (s.get("source_image") and (paper.parent / s["source_image"]).exists())]
                  for k in ("figure", "table")}
    n_assets = Counter("table" if a["kind"] == "table" else "figure" for a in merged)
    for a in merged:
        akind = "table" if a["kind"] == "table" else "figure"
        cands = []
        for s in slots:
            sc, why = score(a, s)
            if sc > 0 and n_assets[akind] == 1 and len(open_slots[akind]) == 1 and open_slots[akind][0] is s:
                sc += 25
                why.append(f"only supplied {akind} and only open {akind} slot")
            if sc > 0:
                cands.append({"slot": s["key"], "score": sc, "why": why})
        cands.sort(key=lambda c: -c["score"])
        a["candidates"] = cands[:4]
        best = cands[0] if cands else None
        second = cands[1]["score"] if len(cands) > 1 else 0
        if not best or best["score"] < 30:
            a["match"] = "unmatched"
        elif best["score"] - second < 15:
            a["match"] = "ambiguous"
        else:
            a["match"] = "matched"
            a["proposed_slot"] = best["slot"]

    # slot competition
    by_slot = {}
    for a in merged:
        if a.get("proposed_slot"):
            by_slot.setdefault(a["proposed_slot"], []).append(a)
    for key, lst in by_slot.items():
        if len(lst) > 1:
            for a in lst:
                a["match"] = "conflict"
            issues.append({"severity": "ASK", "msg": f"{len(lst)} supplied assets compete for {key}: " +
                           ", ".join(Path(a['path']).name for a in lst)})

    for a in merged:
        name = Path(a["path"]).name
        if a["match"] == "unmatched":
            issues.append({"severity": "ASK", "msg": f"{name}: no manuscript figure/table matches; ask whether to include (where, with which caption) or leave out"})
        elif a["match"] == "ambiguous":
            issues.append({"severity": "ASK", "msg": f"{name}: placement ambiguous between " +
                           ", ".join(f"{c['slot']}({c['score']})" for c in a["candidates"][:3])})
        if a.get("unsupported"):
            issues.append({"severity": "ASK", "msg": f"{name}: {a['unsupported']}"})
        if a.get("table_as_image"):
            issues.append({"severity": "ASK", "msg": f"{name}: looks like a table supplied as an image; tables must be editable - ask for the data (CSV/LaTeX) or approval to include it as a figure"})
        if a.get("pages") and a["pages"] > 1:
            issues.append({"severity": "ASK", "msg": f"{name}: PDF has {a['pages']} pages; which page is the figure?"})
        if a.get("fonts_not_embedded"):
            issues.append({"severity": "WARN", "msg": f"{name}: fonts not embedded {a['fonts_not_embedded']} (may render differently; do not alter the file - report)"})
        if a.get("has_foreign_object"):
            issues.append({"severity": "WARN", "msg": f"{name}: SVG uses <foreignObject> (HTML text) - most converters drop it; verify the converted PDF visually"})
        if a.get("external_refs"):
            issues.append({"severity": "WARN", "msg": f"{name}: SVG references external files {a['external_refs'][:3]}"})
        if not a["vector"] if "vector" in a else False:
            dpi = a.get("dpi")
            if a.get("px_w") and a["px_w"] < 1000:
                issues.append({"severity": "WARN", "msg": f"{name}: raster only {a['px_w']}px wide; at column width that is ~{round(a['px_w'] / 3.4)} dpi (<300). Ask for a vector or higher-resolution version if available"})
            if dpi is None:
                issues.append({"severity": "INFO", "msg": f"{name}: no DPI metadata; always set an explicit width in \\includegraphics"})
        if len(a["variants"]) > 1:
            issues.append({"severity": "INFO", "msg": f"{a['group']}: variants {[Path(v).suffix for v in a['variants']]}; preferred {a['ext']} (vector first)"})
        if a.get("sidecar_caption") and a.get("proposed_slot"):
            s = next(s for s in slots if s["key"] == a["proposed_slot"])
            if s.get("caption") and alnum(s["caption"]) != alnum(a["sidecar_caption"]):
                issues.append({"severity": "ASK", "msg": f"{name}: sidecar caption differs from the manuscript caption of {s['key']}; which is authoritative?"})
        if a["kind"] == "table" and a.get("proposed_slot"):
            s = next(s for s in slots if s["key"] == a["proposed_slot"])
            if s.get("manuscript_table"):
                d = table_discrepancy(a["rows"], s["manuscript_table"])
                a["manuscript_comparison"] = d
                if d.get("differences") or d.get("only_in_manuscript") or d.get("only_in_supplied"):
                    issues.append({"severity": "DISCREPANCY", "msg": f"{name} vs manuscript {s['key']}: {json.dumps(d)[:600]}"})
        if a["kind"] == "figure" and a.get("proposed_slot"):
            s = next(s for s in slots if s["key"] == a["proposed_slot"])
            if s.get("source_image"):
                src = (paper.parent / s["source_image"])
                if src.exists() and sha256(src) != a["sha256"]:
                    issues.append({"severity": "VERIFY", "msg": f"{name}: differs from the manuscript's own image {s['source_image']} for {s['key']} (e.g. vector version of the same plot?). Compare renders; if content differs, ask which is authoritative"})
    for h, ps in hashes.items():
        if len(ps) > 1:
            issues.append({"severity": "INFO", "msg": f"identical files supplied: {[Path(p).name for p in ps]}"})
    claimed = {a.get("proposed_slot") for a in merged}
    for s in slots:
        has_src = s.get("source_image") and (paper.parent / s["source_image"]).exists()
        if s["key"] not in claimed and not has_src and not (s["kind"] == "table" and s.get("manuscript_table")):
            if s["caption"] is not None or s["placeholder"]:
                issues.append({"severity": "ASK", "msg": f"manuscript {s['key']} (caption: {str(s['caption'])[:60]!r}) has no image/table in the source and no supplied asset matches"})
    for o in others:
        if o["class"] in ("document", "unsupported"):
            issues.append({"severity": "INFO", "msg": f"not a visual, not included: {Path(o['path']).name} ({o['class']})"})

    manifest = {"paper": str(paper.resolve()), "supplied_paths": [str(Path(p)) for p in args[2:]],
                "assets": merged, "slots": slots, "other_files": others, "issues": issues}
    out = project / "_conversion" / "visuals_manifest.json"
    out.write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")

    print(f"supplied files: {len(files)} | visual assets: {len(merged)} | other files: {len(others)} | manuscript slots: {len(slots)}")
    for s in slots:
        print(f"  slot {s['key']:10s} mentions={s['mentions']} caption={str(s['caption'])[:60]!r} "
              f"src_image={s['source_image']} table_in_ms={'yes' if s['manuscript_table'] else 'no'}")
    for a in merged:
        best = a["candidates"][0] if a["candidates"] else None
        print(f"  asset {Path(a['path']).name:32s} {a['kind']:6s} {a['ext']:5s} vector={a.get('vector', '-')!s:5s} "
              f"match={a['match']:9s} -> {best['slot'] if best else '-'} ({best['score'] if best else 0}) "
              f"{'; '.join(best['why']) if best else ''}")
    for i in issues:
        print(f"  [{i['severity']}] {i['msg']}")
    print(f"manifest: {out}")


# --------------------------------------------------------------- convert
def svg_converters():
    conv = []
    for exe in ("inkscape", "rsvg-convert"):
        if shutil.which(exe):
            conv.append(exe)
    for mod in ("cairosvg", "svglib"):
        try:
            __import__(mod)
            conv.append(mod)
        except Exception:  # noqa: BLE001
            pass
    return conv


def convert_file(src, dest_dir, name=None):
    src, dest_dir = Path(src), Path(dest_dir)
    dest_dir.mkdir(parents=True, exist_ok=True)
    ext = src.suffix.lower()
    stem = safe_stem(name or src.stem)
    rec = {"source": str(src.resolve()), "source_sha256": sha256(src)}
    if ext in (".pdf", ".png", ".jpg", ".jpeg"):
        out = dest_dir / f"{stem}{'.jpg' if ext == '.jpeg' else ext}"
        shutil.copy2(src, out)
        rec.update({"method": "copy (unchanged)", "output": str(out)})
    elif ext == ".eps":
        out = dest_dir / f"{stem}.pdf"
        if not shutil.which("epstopdf"):
            return dict(rec, error="epstopdf not available"), 3
        rc, o, e = run(["epstopdf", f"--outfile={out}", str(src)], timeout=300)
        if rc != 0 or not out.exists():
            return dict(rec, error=f"epstopdf failed: {e[-400:]}"), 1
        rec.update({"method": "epstopdf", "output": str(out)})
    elif ext == ".svg":
        out = dest_dir / f"{stem}.pdf"
        convs = svg_converters()
        if not convs:
            return dict(rec, error=("no SVG converter available (inkscape, rsvg-convert, cairosvg, svglib). "
                                    "Options: (a) user supplies a PDF/PNG export of the same figure; "
                                    "(b) user installs Inkscape or `pip install svglib` / `pip install cairosvg`; "
                                    "(c) keep the SVG and use the LaTeX `svg` package (needs Inkscape + shell-escape; "
                                    "works on Overleaf, not on this machine). Never rasterise via screenshots.")), 3
        method, err = None, ""
        for c in convs:
            try:
                if c == "inkscape":
                    rc, _, err = run(["inkscape", str(src), "--export-type=pdf", f"--export-filename={out}", "--export-text-to-path=false"], 300)
                elif c == "rsvg-convert":
                    rc, _, err = run(["rsvg-convert", "-f", "pdf", "-o", str(out), str(src)], 300)
                elif c == "cairosvg":
                    import cairosvg  # type: ignore
                    cairosvg.svg2pdf(url=str(src), write_to=str(out))
                    rc = 0
                else:
                    from svglib.svglib import svg2rlg  # type: ignore
                    from reportlab.graphics import renderPDF  # type: ignore
                    renderPDF.drawToFile(svg2rlg(str(src)), str(out))
                    rc = 0
                if rc == 0 and out.exists():
                    method = c
                    break
            except Exception as exc:  # noqa: BLE001
                err = str(exc)
        if not method:
            return dict(rec, error=f"SVG conversion failed with {convs}: {err[-400:]}"), 1
        rec.update({"method": f"svg->pdf via {method}", "output": str(out)})
        si, pi = svg_info(src), pdf_info(out)
        missing = sorted(tokens(si.get("text", "")) - tokens(pi.get("text", "")))
        rec["verify"] = {"svg_text_tokens_missing_in_pdf": missing}
        vb = [float(x) for x in re.split(r"[\s,]+", si.get("viewBox", "").strip())] if si.get("viewBox") else None
        if vb and len(vb) == 4 and pi.get("width_pt"):
            r1, r2 = vb[2] / vb[3], pi["width_pt"] / pi["height_pt"]
            rec["verify"]["aspect_ratio_svg_pdf"] = [round(r1, 4), round(r2, 4)]
            rec["verify"]["aspect_ok"] = abs(r1 - r2) / r1 < 0.01
    elif ext in (".tif", ".tiff", ".gif", ".bmp", ".webp"):
        out = dest_dir / f"{stem}.png"
        try:
            from PIL import Image  # type: ignore
            with Image.open(src) as im:
                if getattr(im, "n_frames", 1) > 1:
                    return dict(rec, error=f"{ext} has {im.n_frames} frames; ask which frame"), 1
                im.save(out, "PNG", dpi=im.info.get("dpi", (300, 300)))
        except ImportError:
            return dict(rec, error="Pillow not available for raster conversion"), 3
        rec.update({"method": f"{ext[1:]}->png (Pillow, lossless)", "output": str(out)})
    else:
        return dict(rec, error=f"unsupported figure format {ext}"), 1
    rec["output_sha256"] = sha256(out)
    rec["output_name"] = out.name
    return rec, 0


def cmd_convert(args):
    if len(args) < 2:
        sys.exit("usage: visuals.py convert <asset_file> <dest_dir> [--name <stem>]")
    name = args[args.index("--name") + 1] if "--name" in args else None
    rec, rc = convert_file(args[0], args[1], name)
    print(json.dumps(rec, indent=2, ensure_ascii=False))
    sys.exit(rc)


# ----------------------------------------------------------------- table
UNI_MATH = {"\u00b1": r"$\pm$", "\u2212": r"$-$", "\u00d7": r"$\times$", "\u2264": r"$\le$", "\u2265": r"$\ge$",
            "\u0394": r"$\Delta$", "\u03b1": r"$\alpha$", "\u03b2": r"$\beta$", "\u03b3": r"$\gamma$",
            "\u03b4": r"$\delta$", "\u03b5": r"$\epsilon$", "\u03bc": r"$\mu$", "\u03c3": r"$\sigma$",
            "\u03c1": r"$\rho$", "\u03c4": r"$\tau$", "\u03bb": r"$\lambda$", "\u2192": r"$\rightarrow$",
            "\u2248": r"$\approx$", "\u2020": r"\textsuperscript{\dag}", "\u2021": r"\textsuperscript{\ddag}",
            "\u2013": "--", "\u2014": "---"}


def esc_text(s):
    out = []
    for ch in s:
        if ch in "&%#_{}":
            out.append("\\" + ch)
        elif ch == "~":
            out.append(r"\textasciitilde{}")
        elif ch == "^":
            out.append(r"\textasciicircum{}")
        elif ch == "\\":
            out.append(r"\textbackslash{}")
        elif ch == "<":
            out.append(r"$<$")
        elif ch == ">":
            out.append(r"$>$")
        elif ch in UNI_MATH:
            out.append(UNI_MATH[ch])
        else:
            out.append(ch)
    return "".join(out)


def cell_to_tex(cell, markdown):
    parts = re.split(r"(?<!\\)(\$[^$]*\$)", cell)  # keep inline math verbatim
    res = []
    for k, part in enumerate(parts):
        if part.startswith("$") and part.endswith("$") and len(part) > 1:
            res.append(part)
            continue
        if markdown:
            segs = re.split(r"(\*\*[^*]+\*\*|__[^_]+__|\*[^*]+\*|`[^`]+`)", part)
            for sgm in segs:
                if re.fullmatch(r"\*\*[^*]+\*\*|__[^_]+__", sgm):
                    res.append(r"\textbf{" + esc_text(sgm[2:-2]) + "}")
                elif re.fullmatch(r"\*[^*]+\*", sgm):
                    res.append(r"\emph{" + esc_text(sgm[1:-1]) + "}")
                elif re.fullmatch(r"`[^`]+`", sgm):
                    res.append(r"\texttt{" + esc_text(sgm[1:-1]) + "}")
                else:
                    res.append(esc_text(sgm))
        else:
            res.append(esc_text(part))
    return "".join(res).strip()


def is_numeric_cell(c):
    c = re.sub(r"[*_`$]", "", c).strip()
    return bool(re.match(r"^[\[(]?\s*[<>\u2264\u2265~\u2248+\-\u2212]?\s*\d", c))


def cmd_table(args):
    if len(args) < 2:
        sys.exit("usage: visuals.py table <table_file> <out_tex> [--caption-file F | --caption TEXT] [--label L] [--wide auto|yes|no] [--placement t]")
    src, out = Path(args[0]), Path(args[1])
    opt = lambda k, d=None: args[args.index(k) + 1] if k in args else d  # noqa: E731
    out.parent.mkdir(parents=True, exist_ok=True)
    if src.suffix.lower() == ".tex":
        shutil.copy2(src, out)
        txt = out.read_text(encoding="utf-8", errors="replace")
        print(json.dumps({"method": "copy (unchanged .tex)", "output": str(out), "has_float_env": bool(re.search(r"\\begin\{table\*?\}", txt))}))
        return
    t = read_table(src)
    if t.get("unsupported"):
        sys.exit(t["unsupported"])
    rows = t["rows"]
    if not rows:
        sys.exit("no table rows found")
    ncol = max(len(r) for r in rows)
    rows = [r + [""] * (ncol - len(r)) for r in rows]
    md = src.suffix.lower() in (".md", ".markdown")
    align = t.get("align") if md and t.get("align") and len(t["align"]) == ncol else None
    if not align:
        align = []
        for j in range(ncol):
            body = [r[j] for r in rows[1:] if r[j].strip()]
            align.append("r" if body and sum(is_numeric_cell(c) for c in body) / len(body) > 0.5 else "l")
    caption = opt("--caption")
    if opt("--caption-file"):
        caption = Path(opt("--caption-file")).read_text(encoding="utf-8", errors="replace").strip()
    label = opt("--label")
    est = sum(max(len(re.sub(r"[*_`$\\]", "", r[j])) for r in rows) * 5.0 + 12 for j in range(ncol))
    wide = opt("--wide", "auto")
    star = "*" if wide == "yes" or (wide == "auto" and est > 240) else ""
    lines = [f"\\begin{{table{star}}}[{opt('--placement', 't')}]", "\\centering"]
    if caption:
        lines.append("\\caption{" + caption + "}" + (f"\\label{{{label}}}" if label else ""))
    else:
        lines.append("% NO CAPTION SUPPLIED: the manuscript/sidecar has none - do not invent one (ask the user)")
    lines += [f"\\begin{{tabular}}{{@{{}}{''.join(align)}@{{}}}}", "\\toprule"]
    for i, r in enumerate(rows):
        lines.append(" & ".join(cell_to_tex(c, md) for c in r) + " \\\\")
        if i == 0:
            lines.append("\\midrule")
    lines += ["\\bottomrule", "\\end{tabular}", f"\\end{{table{star}}}", ""]
    out.write_text("\n".join(lines), encoding="utf-8", newline="\n")
    src_nums = Counter(numbers(" ".join(" ".join(r) for r in rows)))
    out_nums = Counter(tex_table_numbers(out.read_text(encoding="utf-8")))
    lost, added = sorted((src_nums - out_nums).elements()), sorted((out_nums - src_nums).elements())
    print(json.dumps({"method": f"{src.suffix[1:]} -> booktabs table{star}", "output": str(out), "shape": [len(rows), ncol],
                      "estimated_width_pt": round(est), "numbers_lost": lost, "numbers_added": added}, ensure_ascii=False))
    if lost or added:
        sys.exit("number mismatch after conversion - inspect the table file")


# ----------------------------------------------------------------- check
def label_page(aux, label):
    m = re.search(r"\\newlabel\{" + re.escape(label) + r"\}\{\{([^{}]*)\}\{([^{}]*)\}", aux)
    if not m:
        return None, None
    try:
        return m.group(1), int(m.group(2))
    except ValueError:
        return m.group(1), None


def load_tex_all(root):
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from lint_project import load_tex_tree  # noqa: E402
    files = load_tex_tree(root, "main.tex", [])
    return files


def float_for_label(files, label):
    """Return (relpath, float_text, env_name) for the float that defines label."""
    for rel, text in files:
        i = text.find("\\label{" + label + "}")
        if i < 0:
            continue
        envs = list(re.finditer(r"\\begin\{(figure|table)(\*?)\}", text[:i]))
        if not envs:
            return rel, text, None
        b = envs[-1]
        e = text.find("\\end{" + b.group(1) + b.group(2) + "}", i)
        return rel, text[b.start():e if e > 0 else len(text)], b.group(1) + b.group(2)
    return None, None, None


SYMBOL_NAMES = {"±": "pm", "×": "times", "≤": "le", "≥": "ge", "Δ": "Delta",
                "α": "alpha", "β": "beta", "γ": "gamma", "σ": "sigma", "μ": "mu",
                "·": "cdot", "−": "-"}


def caption_key(text, is_tex):
    """Comparable form of a caption: TeX captions are de-TeXed; plain captions keep '%' as text."""
    t = tex_to_plain(text) if is_tex else text
    t = t.replace("\\leq", "le").replace("\\geq", "ge")
    for ch, name in SYMBOL_NAMES.items():
        t = t.replace(ch, f" {name} ")
    return alnum(t)


def caption_of(block):
    m = re.search(r"\\caption(?:\[[^\]]*\])?\s*(?=\{)", block or "")
    if not m:
        return None
    c, _ = brace_arg(block, m.end())
    return c


def parse_len(spec, colw, textw, env):
    m = re.match(r"\s*([\d.]*)\s*\\(columnwidth|linewidth|textwidth|hsize)", spec)
    if m:
        f = float(m.group(1) or 1)
        base = {"columnwidth": colw, "textwidth": textw, "hsize": colw}.get(m.group(2))
        if m.group(2) == "linewidth":
            base = textw if env and env.endswith("*") else colw
        return f * base if base else None
    m = re.match(r"\s*([\d.]+)\s*(pt|in|cm|mm|bp)", spec)
    if m:
        return float(m.group(1)) * {"pt": 1, "in": 72.27, "cm": 28.4528, "mm": 2.84528, "bp": 1.00375}[m.group(2)]
    return None


def edge_touch(png_path):
    try:
        from PIL import Image  # type: ignore
    except ImportError:
        return None
    with Image.open(png_path) as im:
        g = im.convert("L")
        w, h = g.size
        px = g.load()
        border = [px[x, 0] for x in range(w)] + [px[x, h - 1] for x in range(w)] + \
                 [px[0, y] for y in range(h)] + [px[w - 1, y] for y in range(h)]
    dark = sum(1 for v in border if v < 200)
    return round(dark / max(1, len(border)), 4)


def render_pdf_page(pdf, page, out_png_stem, dpi):
    run(["pdftoppm", "-r", str(dpi), "-f", str(page), "-l", str(page), "-singlefile", "-png", str(pdf), str(out_png_stem)], 300)
    return Path(str(out_png_stem) + ".png")


def cmd_check(args):
    if not args:
        sys.exit("usage: visuals.py check <project_dir> [--manifest M] [--render-dpi 110]")
    root = Path(args[0]).resolve()
    mpath = Path(args[args.index("--manifest") + 1]) if "--manifest" in args else root / "_conversion" / "visuals_manifest.json"
    dpi = int(args[args.index("--render-dpi") + 1]) if "--render-dpi" in args else 110
    if not mpath.exists():
        sys.exit(f"manifest not found: {mpath} (run `visuals.py inventory` first)")
    man = json.loads(mpath.read_text(encoding="utf-8"))
    files = load_tex_all(root)
    alltex = "\n".join(t for _, t in files)
    aux_p, log_p, pdf_p = root / "_build" / "main.aux", root / "_build" / "main.log", root / "main.pdf"
    aux = aux_p.read_text(encoding="utf-8", errors="replace") if aux_p.exists() else ""
    log = log_p.read_text(encoding="utf-8", errors="replace") if log_p.exists() else ""
    m = re.search(r"PL-LAYOUT columnwidth=([\d.]+)pt textwidth=([\d.]+)pt", log)
    colw, textw = (float(m.group(1)), float(m.group(2))) if m else (None, None)
    _, refs_page = label_page(aux, "pl:refsstart")
    rdir = root / "_conversion" / "render"
    rdir.mkdir(parents=True, exist_ok=True)
    results, errors, warns, asks = [], [], [], []

    def E(msg):
        errors.append(msg)

    if not aux:
        E("no _build/main.aux: compile first (compile_and_check.py)")
    if colw is None:
        warns.append("PL-LAYOUT line not in log (template \\typeout missing): legibility checks use no scaling info")

    for a in man["assets"]:
        name = Path(a["path"]).name
        d = a.get("decision") or {}
        r = {"asset": name, "action": d.get("action"), "slot": d.get("slot"), "label": d.get("label")}
        act = d.get("action")
        if act not in ("include", "exclude", "ask"):
            E(f"{name}: no decision recorded (include / exclude / ask) - supplied materials may not be silently omitted")
            results.append(r)
            continue
        if act in ("exclude", "ask"):
            if not d.get("reason"):
                E(f"{name}: '{act}' without a reason")
            (asks if act == "ask" else warns).append(f"{name}: {act.upper()} - {d.get('reason')}")
            results.append(r)
            continue
        pp = d.get("project_path")
        if not pp or not (root / pp).is_file():
            E(f"{name}: project_path missing or not found: {pp}")
            results.append(r)
            continue
        ppath = root / pp
        conv = d.get("conversion")
        if conv:
            if conv.get("output_sha256") != sha256(ppath):
                E(f"{name}: {pp} changed after conversion (sha mismatch)")
            if conv.get("source_sha256") != a["sha256"]:
                E(f"{name}: conversion record is for a different source file")
            v = conv.get("verify") or {}
            if v.get("svg_text_tokens_missing_in_pdf"):
                E(f"{name}: converted PDF lost text labels {v['svg_text_tokens_missing_in_pdf'][:10]}")
            if v.get("aspect_ok") is False:
                E(f"{name}: converted PDF aspect ratio differs from SVG {v.get('aspect_ratio_svg_pdf')}")
            r["file_identity"] = conv.get("method")
        elif a["kind"] != "table":
            if sha256(ppath) != a["sha256"]:
                E(f"{name}: {pp} is not byte-identical to the supplied file and has no conversion record")
            r["file_identity"] = "byte-identical copy"
        label = d.get("label")
        rel, block, env = float_for_label(files, label) if label else (None, None, None)
        if a["kind"] == "table":
            rel_in = pp[:-4] if pp.endswith(".tex") else pp
            if not re.search(r"\\input\{" + re.escape(rel_in) + r"(\.tex)?\}", alltex):
                E(f"{name}: {pp} is never \\input")
            body_nums = Counter(tex_table_numbers(ppath.read_text(encoding="utf-8", errors="replace")))
            src_nums = Counter(a.get("numbers") or [])
            lost, added = sorted((src_nums - body_nums).elements()), sorted((body_nums - src_nums).elements())
            r["table_numbers"] = {"supplied": sum(src_nums.values()), "in_project": sum(body_nums.values()),
                                  "lost": lost, "added": added}
            if lost or added:
                E(f"{name}: table numbers differ from the supplied file: lost {lost[:10]} added {added[:10]}")
        else:
            finfo = pdf_info(ppath) if ppath.suffix.lower() == ".pdf" else raster_info(ppath)
            gp = [""] + re.findall(r"\{([^{}]*)\}", "".join(re.findall(r"\\graphicspath\s*\{((?:\{[^{}]*\})+)\}", alltex)))
            hits = [m for m in re.finditer(r"\\includegraphics\*?(\[[^\]]*\])?\{([^}]*)\}", alltex)
                    if any(str(Path(g + m.group(2)).with_suffix("")).replace("\\", "/") == str(Path(pp).with_suffix("")).replace("\\", "/")
                           for g in gp)]
            if not hits:
                E(f"{name}: {pp} is never \\includegraphics'd")
            for h in hits:
                opts = h.group(1) or ""
                for bad in ("trim", "clip", "viewport", "bb="):
                    if bad in opts:
                        E(f"{name}: \\includegraphics uses '{bad}' (crops content) - not allowed without user approval")
                if "angle" in opts:
                    warns.append(f"{name}: rotated with angle=; confirm the original orientation is wrong")
                if "width" in opts and "height" in opts and "keepaspectratio" not in opts:
                    E(f"{name}: both width and height set without keepaspectratio (distorts the figure)")
                mw = re.search(r"width\s*=\s*([^,\]]+)", opts)
                ms = re.search(r"scale\s*=\s*([\d.]+)", opts)
                rendered = parse_len(mw.group(1), colw, textw, env) if mw else None
                natural = finfo.get("width_pt")
                if finfo.get("px_w") and not natural:
                    natural = finfo["px_w"] / (finfo.get("dpi") or 72) * 72.27
                if ms and natural:
                    rendered = float(ms.group(1)) * natural
                if rendered is None and natural and not mw:
                    rendered = natural
                if rendered:
                    r["rendered_width_pt"] = round(rendered, 1)
                    limit = textw if env and env.endswith("*") else colw
                    if limit and rendered > limit + 1:
                        E(f"{name}: rendered width {rendered:.0f}pt exceeds {'text' if env and env.endswith('*') else 'column'} width {limit:.0f}pt")
                    if natural:
                        sc = rendered / natural
                        r["scale"] = round(sc, 3)
                        if finfo.get("text_height_p10_pt"):
                            eff = finfo["text_height_p10_pt"] * sc
                            r["effective_font_size_pt_p10"] = round(eff, 2)
                            if eff < 4.0:
                                E(f"{name}: text in figure renders at ~{eff:.1f}pt (<4pt) - illegible; use figure* or larger width, or ask")
                            elif eff < 5.5:
                                warns.append(f"{name}: text in figure renders at ~{eff:.1f}pt (small); check legibility in the render")
                    if finfo.get("px_w"):
                        eff_dpi = finfo["px_w"] / (rendered / 72.27)
                        r["effective_dpi"] = round(eff_dpi)
                        if eff_dpi < 150:
                            warns.append(f"{name}: raster at {eff_dpi:.0f} dpi in print (<150); ask for higher resolution/vector if available")
        if not label:
            E(f"{name}: no label recorded")
        elif not block:
            E(f"{name}: \\label{{{label}}} not found inside a figure/table float")
        else:
            cap = caption_of(block)
            want = d.get("caption_text")
            src_kind = d.get("caption_source")
            if src_kind not in ("manuscript", "sidecar", "asset_file", "user", "none"):
                E(f"{name}: caption_source must be manuscript|sidecar|asset_file|user|none (never generated); got {src_kind!r}")
            if want and cap is None:
                E(f"{name}: recorded caption but the float has no \\caption")
            if want and cap is not None and caption_key(cap, True) != caption_key(want, "\\" in want):
                E(f"{name}: caption in LaTeX differs from the recorded {src_kind} caption")
            if not want and cap:
                E(f"{name}: float has a caption but none was recorded as its source (invented caption?)")
            if not want:
                asks.append(f"{name}: no caption available in manuscript or sidecar - ask the user")
            num, page = label_page(aux, label)
            r["number"], r["page"] = num, page
            if d.get("slot"):
                exp = d["slot"].split(":")[1]
                if num and num != exp:
                    E(f"{name}: printed number {num} differs from manuscript {d['slot']}")
            if page and refs_page and d.get("placement", "main") == "main" and page >= refs_page:
                E(f"{name}: placed on page {page}, after the main paper (references start p{refs_page})")
            slot = next((s for s in man["slots"] if s["key"] == d.get("slot")), None)
            if slot and slot.get("mentions") and not re.search(r"\\(?:ref|cref|Cref|autoref)\{" + re.escape(label) + r"\}", alltex):
                E(f"{name}: manuscript mentions {d['slot']} {slot['mentions']}x but \\ref{{{label}}} is never used")
            if page and pdf_p.exists():
                png = render_pdf_page(pdf_p, page, rdir / f"visual_{safe_stem(label)}_p{page}", dpi)
                r["page_render"] = str(png)
                if a["kind"] == "table":
                    rc, ptxt, _ = run(["pdftotext", "-f", str(page), "-l", str(page), str(pdf_p), "-"])
                    pn = Counter(n.lstrip("-") for n in numbers(ptxt))
                    sn = Counter(n.lstrip("-") for n in (a.get("numbers") or []))
                    miss = sorted((sn - pn).elements())
                    r["numbers_missing_on_page"] = miss
                    if miss:
                        warns.append(f"{name}: numbers not extracted from page {page}: {miss[:10]} (check the render; math glyphs may not extract)")
                elif a.get("text"):
                    rc, ptxt, _ = run(["pdftotext", "-f", str(page), "-l", str(page), str(pdf_p), "-"])
                    lost = sorted(tokens(a["text"]) - tokens(ptxt))
                    r["figure_words_missing_on_page"] = lost
                    if lost:
                        warns.append(f"{name}: figure words not found on page {page}: {lost[:10]} (possible clipping; inspect render)")
        # render the supplied asset itself for side-by-side comparison
        if a["kind"] == "figure":
            src_for_render = ppath if ppath.suffix.lower() == ".pdf" else Path(a["path"])
            if src_for_render.suffix.lower() == ".pdf":
                png = render_pdf_page(src_for_render, 1, rdir / f"asset_{safe_stem(Path(a['path']).stem)}", dpi)
            elif src_for_render.suffix.lower() in (".png", ".jpg", ".jpeg"):
                png = src_for_render
            else:
                png = None
            if png and png.exists():
                r["asset_render"] = str(png)
                frac = edge_touch(png)
                r["content_touching_edge_fraction"] = frac
                if frac is not None and frac > 0.02:
                    warns.append(f"{name}: content touches the file's edge ({frac:.1%} of border) - may already be cropped in the source; verify visually, do not alter")
        results.append(r)

    # assets present in figures/ or tables/ that the manifest does not cover are fine (manuscript's own)
    report = {"errors": errors, "warnings": warns, "ask_user": asks, "assets": results,
              "layout": {"columnwidth_pt": colw, "textwidth_pt": textw, "refs_start_page": refs_page}}
    (root / "_conversion" / "visuals_check.json").write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    for r in results:
        print("  " + json.dumps({k: v for k, v in r.items() if k not in ("page_render", "asset_render")}, ensure_ascii=False))
    for x in errors:
        print(f"[ERROR] {x}")
    for x in warns:
        print(f"[WARN] {x}")
    for x in asks:
        print(f"[ASK] {x}")
    print(f"renders: {rdir}  (Read every visual_* and asset_* image)")
    print(f"{len(errors)} error(s), {len(warns)} warning(s), {len(asks)} question(s) for the user")
    sys.exit(1 if errors else 0)


def main():
    if len(sys.argv) < 2 or sys.argv[1] not in ("inventory", "convert", "table", "check"):
        print(__doc__)
        sys.exit(2)
    {"inventory": cmd_inventory, "convert": cmd_convert, "table": cmd_table, "check": cmd_check}[sys.argv[1]](sys.argv[2:])


if __name__ == "__main__":
    main()
